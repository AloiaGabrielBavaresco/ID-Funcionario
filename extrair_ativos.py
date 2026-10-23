"""
Extrai os funcionários ATIVOS da Bizneo (ID + nome completo) e salva em Excel.

Alternativa em Python à página web: roda no seu computador, sem Netlify.

Uso:
    pip install requests openpyxl
    # Windows (PowerShell):  $env:BIZNEO_TOKEN="seu-token"
    # Mac/Linux:             export BIZNEO_TOKEN="seu-token"
    python extrair_ativos.py

O token NUNCA fica escrito no código: ele é lido da variável de ambiente
BIZNEO_TOKEN (ou digitado quando o script pedir).
"""

import getpass
import os
import sys
from datetime import date

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DOMINIO = os.environ.get("BIZNEO_DOMAIN", "https://aloiaaerospace.bizneohr.com").rstrip("/")
PAGE_SIZE = 50
MAX_PAGINAS = 200  # trava de segurança contra loop infinito
NOME_ARQUIVO = f"funcionarios-ativos-{date.today().isoformat()}.xlsx"

HEADERS = {"Accept": "application/json", "Content-Type": "application/json"}


def obter_token():
    token = os.environ.get("BIZNEO_TOKEN") or getpass.getpass("Token da API Bizneo: ")
    token = token.strip()
    if not token:
        sys.exit("Token não informado.")
    return token


def buscar_pagina(token, pagina):
    try:
        r = requests.get(
            f"{DOMINIO}/api/v1/users",
            headers=HEADERS,
            params={"token": token, "page": pagina, "page_size": PAGE_SIZE},
            timeout=30,
        )
    except requests.exceptions.RequestException as e:
        # Não imprime a URL (ela contém o token).
        sys.exit(f"Falha de conexão com a Bizneo: {type(e).__name__}")
    if r.status_code in (401, 403):
        sys.exit(f"A Bizneo recusou o acesso (HTTP {r.status_code}). Verifique o token.")
    if r.status_code != 200:
        sys.exit(f"A Bizneo respondeu HTTP {r.status_code}.")
    try:
        return r.json()
    except ValueError:
        sys.exit("Resposta da Bizneo não é um JSON válido.")


def listar_ativos(token):
    ativos = []
    vistos = set()
    for pagina in range(1, MAX_PAGINAS + 1):
        dados = buscar_pagina(token, pagina)
        lista = dados.get("users", []) if isinstance(dados, dict) else []
        if not lista:
            break
        novos = 0
        for u in lista:
            uid = u.get("id")
            if not uid:
                continue
            uid = str(uid)
            if uid in vistos:
                continue
            vistos.add(uid)
            novos += 1
            # Ativo = tem algum contrato de trabalho sem data de fim.
            if any(c.get("end_at") is None for c in u.get("work_contracts", [])):
                nome = f"{(u.get('first_name') or '').strip()} {(u.get('last_name') or '').strip()}".strip()
                ativos.append((uid, nome))
        print(f"  página {pagina}: {len(vistos)} usuários lidos até agora")
        if novos == 0:
            break
    ativos.sort(key=lambda x: x[1].casefold())
    return ativos, len(vistos)


def salvar_excel(ativos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ativos"
    ws.append(["ID", "Nome completo"])
    for linha in ativos:
        ws.append(list(linha))

    for c in (1, 2):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
    larguras = {1: 14, 2: max([len(n) for _, n in ativos] + [14]) + 4}
    for c, w in larguras.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    wb.save(NOME_ARQUIVO)


def main():
    token = obter_token()
    print(f"Consultando {DOMINIO} ...")
    ativos, total = listar_ativos(token)
    print(f"-> {len(ativos)} ativos de {total} usuários.")
    salvar_excel(ativos)
    print(f"Arquivo gerado: {NOME_ARQUIVO}")


if __name__ == "__main__":
    main()
